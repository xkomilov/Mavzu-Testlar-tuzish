import glob,re,sys
sys.path.insert(0,'pylib')
from docx import Document
for f in sorted(glob.glob(sys.argv[1]+'/*.docx')):
    d=Document(f); ps=[p.text for p in d.paragraphs]
    i=[k for k,t in enumerate(ps) if t.strip()=='Foydalanilgan adabiyotlar'][0]
    body=" ".join(ps[1:i]); n=len(body.split())
    bad=re.findall("[’‘ʼ`•]|[\u0400-\u04FF]",body+ps[0])
    nb=sum(1 for p in d.paragraphs if p.style.name=='List Bullet')
    print(f.split('/')[-1],ps[0][:55],n,'bullets',nb,'BAD' if bad else 'ok')
