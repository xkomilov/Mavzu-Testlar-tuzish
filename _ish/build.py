#!/usr/bin/env python3
"""Source .txt -> .docx builder with word-count and format checks.
Markup:  TITLE: ...   |  '## Heading'  |  '- bullet'  |  plain paragraph
Last section '## Foydalanilgan adabiyotlar' = reference list (excluded from count).
"""
import sys, os, re, glob, zipfile
sys.path.insert(0, '/projects/sandbox/_work/pylib')
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

REF_HEAD = 'Foydalanilgan adabiyotlar'


def parse(path):
    title, secs, cur = None, [], None
    for line in open(path, encoding='utf8').read().split('\n'):
        if not line.strip():
            continue
        if line.startswith('TITLE:'):
            title = line[6:].strip()
        elif line.startswith('## '):
            cur = {'h': line[3:].strip(), 'items': []}
            secs.append(cur)
        elif line.startswith('- '):
            cur['items'].append(('b', line[2:].strip()))
        else:
            cur['items'].append(('p', line.strip()))
    return title, secs


def count_words(secs):
    n_all = n_text = 0
    for s in secs:
        if s['h'] == REF_HEAD:
            continue
        hw = len(s['h'].split())
        n_all += hw
        for _, t in s['items']:
            w = len(t.split())
            n_all += w
            n_text += w
    return n_all, n_text


def setfont(run, size=14, bold=False):
    run.font.name = 'Times New Roman'
    run._element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    run.font.size = Pt(size)
    run.bold = bold


def build(path, outdir):
    title, secs = parse(path)
    doc = Document()
    sec = doc.sections[0]
    sec.left_margin = Cm(3); sec.right_margin = Cm(1.5)
    sec.top_margin = Cm(2); sec.bottom_margin = Cm(2)
    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'; st.font.size = Pt(14)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman')
    for sname in ('Title', 'Heading 1', 'Heading 2', 'List Bullet'):
        s = doc.styles[sname]
        s.font.name = 'Times New Roman'; s.font.color.rgb = None
        s.element.rPr.rFonts.set(qn('w:eastAsia'), 'Times New Roman') if s.element.rPr is not None else None
    p = doc.add_paragraph(style='Heading 1')
    r = p.add_run(title); setfont(r, 16, True)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for s in secs:
        hp = doc.add_paragraph(style='Heading 2')
        hr = hp.add_run(s['h']); setfont(hr, 14, True)
        isref = s['h'] == REF_HEAD
        for kind, t in s['items']:
            if kind == 'b':
                bp = doc.add_paragraph(style='List Bullet')
                setfont(bp.add_run(t), 14 if not isref else 12)
                bp.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            else:
                pp = doc.add_paragraph()
                setfont(pp.add_run(t), 12 if isref else 14)
                pp.alignment = WD_ALIGN_PARAGRAPH.LEFT if isref else WD_ALIGN_PARAGRAPH.JUSTIFY
                pp.paragraph_format.first_line_indent = None if isref else Cm(1.25)
                pp.paragraph_format.line_spacing = 1.15
    out = os.path.join(outdir, os.path.basename(path)[:-4] + '.docx')
    doc.save(out)
    return out, title, secs


def check(docx_path, title, secs, lo=800, hi=950):
    problems = []
    n_all, n_text = count_words(secs)
    z = zipfile.ZipFile(docx_path)
    xml = z.read('word/document.xml').decode('utf8')
    plain = re.sub(r'<[^>]+>', '', xml)
    if '\u2019' in plain or '\u02bc' in plain or '\u2018' in plain or '`' in plain:
        problems.append('non-ASCII apostrophe')
    if re.search(r'[\u0400-\u04FF]', plain):
        problems.append('Cyrillic char in text')
    if '•' in plain:
        problems.append('manual bullet char')
    if not (lo <= n_text and n_all <= hi):
        problems.append(f'word count out of range (all={n_all}, text={n_text})')
    heads = [s['h'] for s in secs]
    if not heads or heads[0] != 'Kirish': problems.append('first section not Kirish')
    if 'Xulosa' not in heads: problems.append('no Xulosa')
    if heads[-1] != REF_HEAD: problems.append('no refs at end')
    # citations used must exist in ref list and vice-versa
    body = ' '.join(t for s in secs if s['h'] != REF_HEAD for _, t in s['items'])
    used = set(re.findall(r'\[(\d+)\]', body))
    refs = set()
    for s in secs:
        if s['h'] == REF_HEAD:
            for _, t in s['items']:
                m = re.match(r'\[(\d+)\]', t)
                if m: refs.add(m.group(1))
    if used - refs: problems.append(f'cited but missing in list: {sorted(used-refs)}')
    if refs - used: problems.append(f'listed but never cited: {sorted(refs-used)}')
    return n_all, n_text, problems


if __name__ == '__main__':
    src, outdir = sys.argv[1], sys.argv[2]
    os.makedirs(outdir, exist_ok=True)
    only_check = '--check' in sys.argv
    files = sorted(glob.glob(os.path.join(src, '[0-9][0-9][0-9].txt')))
    bad = 0
    for f in files:
        out, title, secs = build(f, outdir)
        a, t, pr = check(out, title, secs)
        print(os.path.basename(out), f'all={a} text={t}', 'OK' if not pr else 'PROBLEMS: ' + '; '.join(pr))
        bad += bool(pr)
    print('files:', len(files), 'with problems:', bad)
